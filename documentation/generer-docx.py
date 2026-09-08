# -*- coding: utf-8 -*-
"""
Manuel d'administration du site Philippe Grégoire Yacé — version Word.

La charte du site est reprise autant que Word le permet :
  — Bodoni MT pour les titres  (la famille de Bodoni Moda, livrée avec Office)
  — Century Gothic pour le texte (déjà dans la pile de secours de Jost sur le
    site, géométrique comme elle, livrée avec Office)
  — encre #262523, laiton #7D6134, filets #DED7C9, retrait #EFEAE0
  — filets plutôt que cadres, aucune ombre portée

La source de vérité est `manuel-administration.html`, à côté. Ce script en
reprend le contenu pour produire la version Word ; les deux se modifient
ensemble, sans quoi ils divergent.

    python3 -m pip install --target ./pylib python-docx
    PYTHONPATH=./pylib python3 generer-docx.py

Le PDF, lui, se tire directement de la page HTML, qui porte sa feuille
d'impression — sommaire en page de garde, une partie par page :

    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome" \
      --headless=new --no-pdf-header-footer \
      --print-to-pdf=Manuel-administration-site-Yace.pdf \
      "file://$PWD/manuel-administration.html"
"""

import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# --- Charte ---------------------------------------------------------------
ENCRE       = RGBColor(0x26, 0x25, 0x23)
ENCRE2      = RGBColor(0x4D, 0x4D, 0x4D)
ENCRE3      = RGBColor(0x6B, 0x68, 0x62)
LAITON_TXT  = RGBColor(0x7D, 0x61, 0x34)
LAITON      = "A88B5C"
FILET       = "DED7C9"
FILET_FORT  = "C4BBA9"
RETRAIT     = "EFEAE0"
LAVIS       = "F6F1E7"

DISPLAY = "Bodoni MT"
TEXTE   = "Century Gothic"

doc = Document()

# --- Page ------------------------------------------------------------------
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin = Cm(2.4), Cm(2.2)
s.left_margin, s.right_margin = Cm(2.4), Cm(2.4)


def police(style, nom, taille, gras=False, italique=False, couleur=None,
           avant=0, apres=0, interligne=None, espacement=None):
    f = style.font
    f.name = nom
    f.size = Pt(taille)
    f.bold = gras
    f.italic = italique
    if couleur is not None:
        f.color.rgb = couleur
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rf.set(qn(a), nom)
    if espacement is not None:                     # interlettrage, en 1/20 pt
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:val'), str(int(espacement)))
        rpr.append(sp)
    p = style.paragraph_format
    p.space_before = Pt(avant)
    p.space_after = Pt(apres)
    if interligne:
        p.line_spacing = interligne
    return style


# --- Styles ----------------------------------------------------------------
police(doc.styles['Normal'], TEXTE, 10, couleur=ENCRE, apres=7, interligne=1.32)
police(doc.styles['Title'],  DISPLAY, 34, couleur=ENCRE, avant=0, apres=10)
police(doc.styles['Heading 1'], DISPLAY, 22, couleur=ENCRE, avant=0, apres=12)
police(doc.styles['Heading 2'], DISPLAY, 15, couleur=ENCRE, avant=16, apres=6)
police(doc.styles['Heading 3'], TEXTE, 10.5, gras=True, couleur=LAITON_TXT, avant=12, apres=4)
for n in ('Heading 1', 'Heading 2', 'Heading 3'):
    doc.styles[n].font.color.rgb = {'Heading 3': LAITON_TXT}.get(n, ENCRE)

st = doc.styles.add_style('Sur', 1)               # sur-titre en capitales
police(st, TEXTE, 7.5, gras=True, couleur=LAITON_TXT, apres=3, espacement=32)
st = doc.styles.add_style('Lead', 1)
police(st, TEXTE, 11, couleur=ENCRE2, apres=9, interligne=1.35)
st = doc.styles.add_style('Meta', 1)
police(st, TEXTE, 8, couleur=ENCRE3, apres=2)
st = doc.styles.add_style('Puce', 1)
police(st, TEXTE, 10, couleur=ENCRE, apres=4, interligne=1.3)
st.paragraph_format.left_indent = Cm(0.7)
st.paragraph_format.first_line_indent = Cm(-0.35)
st = doc.styles.add_style('Etape', 1)
police(st, TEXTE, 10, couleur=ENCRE, apres=5, interligne=1.3)
st.paragraph_format.left_indent = Cm(0.85)
st.paragraph_format.first_line_indent = Cm(-0.85)
st = doc.styles.add_style('SommPartie', 1)
police(st, TEXTE, 8.5, gras=True, couleur=ENCRE3, avant=9, apres=2, espacement=22)
st = doc.styles.add_style('SommItem', 1)
police(st, TEXTE, 9.5, couleur=ENCRE, apres=1)
st.paragraph_format.left_indent = Cm(0.8)


# --- Petits outils XML -----------------------------------------------------
def bordure(el, cotes, couleur=FILET, taille=4):
    """Pose des filets sur une cellule ou un paragraphe."""
    pr = el.get_or_add_tcPr() if el.tag.endswith('}tc') else el.get_or_add_pPr()
    nom = 'w:tcBorders' if el.tag.endswith('}tc') else 'w:pBdr'
    b = pr.find(qn(nom))
    if b is None:
        b = OxmlElement(nom)
        pr.append(b)
    for cote in cotes:
        e = OxmlElement('w:' + cote)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(taille))
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), couleur)
        b.append(e)


def fond(cell, couleur):
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), couleur)
    cell._tc.get_or_add_tcPr().append(sh)


def marges_cellule(cell, haut=90, bas=90, gauche=0, droite=140):
    mar = OxmlElement('w:tcMar')
    for nom, val in (('top', haut), ('bottom', bas), ('left', gauche), ('right', droite)):
        e = OxmlElement('w:' + nom)
        e.set(qn('w:w'), str(val))
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    cell._tc.get_or_add_tcPr().append(mar)


MOTIF = re.compile(r'(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)')


def ecrire(par, texte, taille=None, couleur=None, nom=None):
    """Écrit un texte avec **gras**, *italique* et `champ`."""
    for morceau in MOTIF.split(texte):
        if not morceau:
            continue
        gras = ital = champ = False
        if morceau.startswith('**'):
            morceau, gras = morceau[2:-2], True
        elif morceau.startswith('`'):
            morceau, champ = morceau[1:-1], True
        elif morceau.startswith('*'):
            morceau, ital = morceau[1:-1], True
        r = par.add_run(morceau)
        r.bold = gras
        r.italic = ital
        r.font.name = nom or TEXTE
        rpr = r._element.get_or_add_rPr()
        rf = OxmlElement('w:rFonts')
        for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
            rf.set(qn(a), nom or TEXTE)
        rpr.append(rf)
        if taille:
            r.font.size = Pt(taille)
        if champ:
            r.bold = True
            r.font.color.rgb = LAITON_TXT
        elif couleur is not None:
            r.font.color.rgb = couleur
    return par


def p(texte, style=None, taille=None, couleur=None):
    par = doc.add_paragraph(style=style)
    ecrire(par, texte, taille=taille, couleur=couleur)
    return par


def puces(items):
    for it in items:
        par = doc.add_paragraph(style='Puce')
        r = par.add_run('—  ')
        r.font.color.rgb = RGBColor(0xA8, 0x8B, 0x5C)
        ecrire(par, it)


def etapes(items):
    for i, it in enumerate(items, 1):
        par = doc.add_paragraph(style='Etape')
        r = par.add_run('%d.  ' % i)
        r.bold = True
        r.font.color.rgb = LAITON_TXT
        ecrire(par, it)


def note(titre, lignes, sourd=False):
    """Encadré : un filet de laiton à gauche, un lavis discret. Pas un cadre."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    c = t.cell(0, 0)
    c.width = Cm(16.2)
    fond(c, RETRAIT if sourd else LAVIS)
    bordure(c._tc, ['left'], FILET_FORT if sourd else LAITON, taille=18)
    marges_cellule(c, haut=140, bas=140, gauche=170, droite=170)
    c.paragraphs[0].text = ''
    par = c.paragraphs[0]
    r = par.add_run(titre)
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.name = TEXTE
    r.font.color.rgb = ENCRE if sourd else LAITON_TXT
    par.paragraph_format.space_after = Pt(4)
    for i, ligne in enumerate(lignes):
        q = c.add_paragraph()
        q.paragraph_format.space_after = Pt(2 if i < len(lignes) - 1 else 0)
        ecrire(q, ligne, taille=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def tableau(entetes, lignes, largeurs=None):
    t = doc.add_table(rows=1, cols=len(entetes))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    for i, e in enumerate(entetes):
        c = t.rows[0].cells[i]
        c.text = ''
        par = c.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        r = par.add_run(e.upper())
        r.bold = True
        r.font.size = Pt(7.5)
        r.font.name = TEXTE
        r.font.color.rgb = ENCRE3
        bordure(c._tc, ['top', 'bottom'], "262523" if i == 0 else "262523", taille=6)
        marges_cellule(c)
    for ligne in lignes:
        cells = t.add_row().cells
        for i, val in enumerate(ligne):
            c = cells[i]
            c.text = ''
            par = c.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            ecrire(par, val, taille=9)
            if i == 0:
                for r in par.runs:
                    r.bold = True
            bordure(c._tc, ['bottom'], FILET, taille=4)
            marges_cellule(c)
    if largeurs:
        for row in t.rows:
            for i, w in enumerate(largeurs):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return t


def fiche(paires):
    """Le résumé en tête d'une rubrique : deux colonnes, deux filets."""
    t = doc.add_table(rows=len(paires), cols=2)
    t.autofit = False
    for i, (cle, val) in enumerate(paires):
        g, d = t.rows[i].cells
        g.width, d.width = Cm(3.4), Cm(12.8)
        g.text = ''
        par = g.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        r = par.add_run(cle.upper())
        r.bold = True
        r.font.size = Pt(7.5)
        r.font.name = TEXTE
        r.font.color.rgb = ENCRE3
        d.text = ''
        par = d.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        ecrire(par, val, taille=9)
        cotes = ['top'] if i == 0 else []
        if i == len(paires) - 1:
            cotes.append('bottom')
        for c in (g, d):
            if cotes:
                bordure(c._tc, cotes, FILET, taille=4)
            marges_cellule(c, haut=70, bas=70)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return t


def partie(num, titre, saut=True):
    par = doc.add_paragraph(style='Sur')
    # Le saut porte sur le paragraphe lui-meme : un saut dans un paragraphe a
    # part laisserait une ligne vide en tete de chaque partie.
    if saut:
        par.paragraph_format.page_break_before = True
    par.add_run('Partie ' + num)
    par.runs[0].font.color.rgb = LAITON_TXT
    h = doc.add_heading(level=1)
    ecrire(h, titre, nom=DISPLAY)
    for r in h.runs:
        r.font.color.rgb = ENCRE
        r.font.size = Pt(22)
    bordure(h._p, ['bottom'], "262523", taille=12)
    h.paragraph_format.space_after = Pt(14)


def rubrique(titre):
    h = doc.add_heading(level=2)
    ecrire(h, titre, nom=DISPLAY)
    for r in h.runs:
        r.font.color.rgb = ENCRE
        r.font.size = Pt(15)


def sous(titre):
    h = doc.add_heading(level=3)
    ecrire(h, titre)
    for r in h.runs:
        r.font.color.rgb = LAITON_TXT
        r.font.size = Pt(10.5)


# =========================================================================
#  COUVERTURE
# =========================================================================
p('Philippe Grégoire Yacé — une destinée', style='Sur')
t = doc.add_paragraph(style='Title')
ecrire(t, 'Administrer le site', nom=DISPLAY)
for r in t.runs:
    r.font.size = Pt(34)
    r.font.color.rgb = ENCRE
bordure(t._p, ['bottom'], FILET, taille=6)
t.paragraph_format.space_after = Pt(14)

p("Tout ce qu'il faut savoir pour tenir le site à jour : publier une actualité, "
  "déposer une photographie, ouvrir la boutique, répondre à un visiteur. "
  "Aucune connaissance technique n'est nécessaire.", style='Lead')

p("Manuel de l'éditeur  ·  Version du 8 septembre 2026", style='Meta')
p("Se lit dans l'ordre, se consulte au besoin.", style='Meta')

doc.add_paragraph()

# --- Sommaire --------------------------------------------------------------
rubrique('Sommaire')
SOMMAIRE = [
    ('01 — Prendre en main', ['Se connecter', 'Qui peut faire quoi',
                              'Brouillon et publié', 'Le tour du propriétaire']),
    ('02 — Les fichiers', ['La médiathèque', 'Formats et poids acceptés',
                           'Dimensions à respecter', 'Préparer une image',
                           'Les vidéos']),
    ('03 — Les contenus', ['Actualités', 'Événements', 'Biographie', 'Repères',
                           'Archives', 'Héritage', 'Points de vente']),
    ('04 — Ce qui vient des visiteurs', ['Témoignages', 'Contributions', 'Messages']),
    ('05 — Le livre et la boutique', ['Paramètres', 'Ouvrir les commandes',
                                      'Zones de livraison', 'Commandes']),
    ('06 — Réglages', ['Comptes', 'Version anglaise']),
    ('07 — Aide-mémoire', ["Ce qui empêche de publier", "Fichiers en un coup d'œil",
                           "Si quelque chose ne marche pas"]),
]
for tit, items in SOMMAIRE:
    p(tit, style='SommPartie')
    for it in items:
        p(it, style='SommItem')

# =========================================================================
#  PARTIE 1
# =========================================================================
partie('01', 'Prendre en main')

rubrique('Se connecter')
p("L'administration du site vit à une adresse séparée : celle du site, suivie de "
  "`/cmsadmin`. Mettez-la en favori, vous y reviendrez tous les jours.")
etapes([
    "Ouvrez cette adresse dans votre navigateur.",
    "Saisissez votre adresse électronique et votre mot de passe.",
    "Vous arrivez sur le tableau de bord. Le menu est à gauche.",
])
note("Après plusieurs essais manqués, l'accès se bloque", [
    "Cinq mots de passe faux à la suite sur le même compte, et la connexion est "
    "refusée pendant un quart d'heure — même avec le bon mot de passe. C'est une "
    "protection contre les intrusions, pas une panne. Attendez, puis réessayez calmement.",
])
p("**Pensez à vous déconnecter** quand vous quittez un ordinateur partagé : le bouton "
  "est en haut à droite. Sinon, la session se ferme d'elle-même au bout d'un long "
  "moment d'inactivité.")

rubrique('Qui peut faire quoi')
p("Il existe deux sortes de comptes. Le menu de gauche n'affiche que ce à quoi votre "
  "compte a droit : si vous ne voyez pas une rubrique décrite ici, c'est que votre "
  "compte n'y a pas accès.")
tableau(['Rôle', "Ce qu'il peut faire"], [
    ['Éditeur', "Tous les contenus, la médiathèque, la modération des témoignages, des "
                "contributions et des messages, les points de vente, les traductions."],
    ['Administrateur', "Tout ce qui précède, **plus** les commandes, les zones de "
                       "livraison et la gestion des comptes."],
], largeurs=[3.6, 12.6])

rubrique('Brouillon et publié — la règle à retenir')
p("C'est le mécanisme le plus important du site, et il vaut pour presque tout ce que "
  "vous saisirez. Chaque contenu porte un état : **Brouillon** ou **Publié**.")
puces([
    "**Brouillon** — le contenu n'existe pas pour les visiteurs. Il n'apparaît sur "
    "aucune page, et même quelqu'un qui devinerait son adresse ne verrait rien. Vous "
    "pouvez le laisser en chantier des semaines.",
    "**Publié** — il est en ligne, tout de suite, pour tout le monde.",
])
p("Un contenu nouveau naît toujours en brouillon. Pour le mettre en ligne, deux chemins : "
  "le menu `Statut` dans sa fiche, ou — plus rapide — le **bouton en forme d'œil** dans "
  "la liste, qui bascule d'un état à l'autre sans ouvrir la fiche.")
note("Dépublier est toujours permis, publier se mérite", [
    "Retirer quelque chose du site ne demande jamais rien. Le mettre en ligne, en "
    "revanche, peut être refusé : il manque une source, une date, une ville. Le site vous "
    "dit alors précisément ce qui manque. La liste complète de ces règles est en partie 07.",
])

rubrique('Le tour du propriétaire')
p("Le menu de gauche est rangé en trois familles.")
tableau(['Famille', 'Rubriques', "Ce qu'on y fait"], [
    ['Contenus',
     "Actualités, Événements, Biographie, Repères, Archives, Héritage, Points de vente, "
     "Médiathèque, Traductions",
     "Tout ce que le visiteur lit et regarde."],
    ['Modération', "Témoignages, Contributions, Messages",
     "Ce que les visiteurs vous envoient. Rien n'est publié sans votre accord."],
    ['Administration', "Zones de livraison, Commandes, Paramètres, Comptes",
     "Le livre, sa vente, et les personnes qui accèdent à l'administration."],
], largeurs=[3.0, 6.2, 7.0])
p("**Toutes les listes fonctionnent pareil.** Un tableau, une barre de recherche, et sur "
  "chaque ligne trois boutons : l'œil pour publier ou dépublier, le crayon pour modifier, "
  "la corbeille pour supprimer. Cliquer sur le titre ouvre la fiche.")
p("**La suppression est définitive.** Il n'y a pas de corbeille de récupération. Le site "
  "demande confirmation, et c'est la seule barrière : en cas de doute, dépubliez plutôt "
  "que de supprimer.")

# =========================================================================
#  PARTIE 2
# =========================================================================
partie('02', 'Les fichiers')

rubrique('La médiathèque : le magasin de tous les fichiers')
p("Photographies, documents scannés, enregistrements sonores : **tout passe par la "
  "médiathèque**, et une seule fois. On ne dépose jamais un fichier depuis la fiche d'une "
  "actualité ou d'une archive — on l'y rattache, en le choisissant dans la médiathèque.")
p("L'ordre de travail est donc toujours le même :")
etapes([
    "**Déposer** le fichier dans la médiathèque.",
    "**Le renseigner** : titre, légende, crédit, catégorie.",
    "**Le rattacher** depuis la fiche qui en a besoin.",
])
p("Un même fichier peut servir à plusieurs endroits. Le déposer deux fois ne sert à rien "
  "et encombre le serveur.")

sous('Déposer des fichiers')
p("Le bouton de dépôt est en haut de l'écran Médiathèque. Vous pouvez en sélectionner "
  "plusieurs d'un coup — **vingt au maximum par envoi**. Au-delà, faites deux voyages.")
p("Le nom de vos fichiers n'a aucune importance : le site le réécrit automatiquement. Ce "
  "qui compte, c'est ce que vous saisissez ensuite.")

sous('Renseigner un fichier')
tableau(['Champ', 'À quoi il sert'], [
    ['Titre', "Obligatoire. Vous permet de retrouver le fichier dans la médiathèque. Soyez "
              "descriptif : « Yacé à l'Assemblée, 1962 » et non « photo 3 »."],
    ['Légende', "Le texte affiché sous l'image sur le site. **Sert aussi de description aux "
                "personnes non voyantes et aux moteurs de recherche** — écrivez-la, elle "
                "compte double."],
    ['Crédit', "Le photographe, le fonds ou l'institution qui détient la pièce. "
               "**Obligatoire pour publier une image.**"],
    ['Date de prise de vue', "Texte libre : « 1962 », « vers 1958 », « mars 1970 »."],
    ['Catégorie', "Portrait, Officiel, Privé, Document, Presse. Sert à filtrer la "
                  "médiathèque quand elle sera grande."],
    ['Rang', "L'ordre d'apparition dans la galerie. Le plus petit passe en premier."],
], largeurs=[4.0, 12.2])
note("Une image publiée doit porter son crédit", [
    "Le site refuse de publier une image dont le champ Crédit est vide. Ce n'est pas une "
    "formalité : le fonds recevra des documents prêtés par des familles et des organes de "
    "presse, et savoir de qui vient quoi protège l'éditeur.",
])

rubrique('Formats et poids acceptés')
p("Trois familles de fichiers, trois limites de poids différentes. Un scan de vingt pages "
  "n'a pas le même poids naturel qu'une photographie.")
tableau(['Famille', 'Formats acceptés', 'Poids maximum', 'Pour quoi'], [
    ['Images', 'JPEG (.jpg), PNG, WebP', '8 Mo',
     'Photographies, portraits, reproductions'],
    ['Documents', 'PDF uniquement', '30 Mo',
     'Correspondances, discours dactylographiés, coupures scannées'],
    ['Enregistrements', 'MP3, M4A, OGG', '60 Mo',
     'Discours, allocutions, entretiens'],
], largeurs=[3.2, 3.9, 2.6, 6.5])
p("**Deux limites supplémentaires sur les images :**")
puces([
    "**40 millions de pixels au maximum** — soit, par exemple, 8 000 × 5 000 points. C'est "
    "très largement au-dessus de ce qu'un appareil photo courant produit ; la limite ne se "
    "rencontre qu'avec un scan de très grand format.",
    "**Le format est vérifié dans le fichier lui-même**, pas dans son nom. Renommer un "
    "document en `.jpg` ne le fera pas passer pour une image — le site le refusera, et "
    "c'est voulu.",
])
note("Le site fabrique tout seul les versions réduites", [
    "À chaque dépôt d'image, deux copies allégées sont créées automatiquement : une petite "
    "pour les vignettes, une moyenne pour l'affichage courant. Vous n'avez rien à préparer "
    "de ce côté-là. Les photographies prises de travers par un téléphone sont également "
    "redressées toutes seules.",
], sourd=True)

rubrique('Dimensions à respecter')
p("Quelques emplacements du site ont une taille et un cadrage précis, parce que l'image y "
  "occupe une place fixe dans la maquette. Partout ailleurs, le site recadre tout seul.")

sous('Les emplacements à dimensions imposées')
tableau(['Emplacement', 'Dimensions', 'Cadrage attendu'], [
    ["Grande image d'accueil (trois images qui défilent)", '2000 × 2600 px',
     "Portrait. **Placez le sujet dans la moitié droite** : la moitié gauche est recouverte "
     "par un voile clair sur lequel s'écrit le texte."],
    ['Couverture du livre', '1200 × 1550 px',
     "La couverture seule, sans décor autour. Proportion du livre réel (240 × 310 mm)."],
    ['Portrait en tête de biographie', '1400 × 1750 px',
     "Buste, cadrage serré. *Déjà livré.*"],
    ["Portrait de l'auteur du livre", '1000 × 1250 px',
     "Portrait vertical classique."],
    ['Extraits du livre', '1500 × 1000 px',
     "Double page photographiée à plat, horizontale."],
], largeurs=[4.6, 3.0, 8.6])

sous('Partout ailleurs : aucune dimension imposée')
p("Les archives, la galerie, les illustrations d'actualités, les images d'Héritage et de "
  "la biographie sont recadrées automatiquement selon quatre proportions qui alternent "
  "dans la page : carré, portrait, paysage et panoramique.")
p("**La seule règle : le sujet doit tenir au centre.** Une photographie cadrée très serré "
  "sur un bord sera coupée sur au moins une des quatre proportions. Laissez un peu d'air "
  "autour du sujet.")
note("Les archives sont uniformisées automatiquement", [
    "Le site applique de lui-même un traitement noir et blanc chaud aux images d'archives, "
    "ce qui donne une allure homogène à des sources d'origines très diverses. **Ne "
    "retouchez pas vos scans** : déposez-les tels quels, en couleur si c'est ainsi qu'ils "
    "vous parviennent.",
], sourd=True)

rubrique('Préparer une image avant de la déposer')
p("Les fichiers d'origine sont presque toujours bien plus lourds que nécessaire. Une "
  "minute de préparation fait gagner beaucoup de temps de chargement aux visiteurs — dont "
  "beaucoup consultent le site depuis un téléphone.")
etapes([
    "**Redimensionnez** à la taille du tableau ci-dessus si l'emplacement en impose une. "
    "Sinon, visez **1600 points sur le plus grand côté** : c'est la taille que le site "
    "utilise pour l'affichage courant, au-delà le poids est perdu.",
    "**Enregistrez en JPEG, qualité 80.** C'est le réglage qui donne le meilleur rapport "
    "entre le poids et la netteté. La différence avec la qualité 100 ne se voit pas à "
    "l'écran ; le poids, lui, est divisé par trois ou quatre.",
    "**Vérifiez le poids final.** Une photographie bien préparée pèse entre 150 et 600 Ko. "
    "Si vous approchez des 8 Mo, quelque chose n'a pas été redimensionné.",
])
p("Le PNG n'a d'intérêt que pour un logo ou un dessin au trait sur fond transparent. Pour "
  "une photographie, il pèse cinq fois plus qu'un JPEG sans rien apporter.")

rubrique('Les vidéos ne se déposent pas')
p("Une heure d'archive vidéo pèse plusieurs gigaoctets et saturerait le site dès la "
  "première consultation groupée. Les vidéos restent donc chez leur hébergeur.")
p("La marche à suivre : mettez la vidéo en ligne sur **YouTube**, puis collez son adresse "
  "dans le champ `Lien de la vidéo` de la fiche d'archive ou d'Héritage. Le site "
  "l'affiche dans la page.")
p("Seules les adresses YouTube sont reconnues — les deux formes `youtube.com/watch…` et "
  "`youtu.be/…` fonctionnent. Une adresse d'une autre plateforme sera refusée à la saisie, "
  "plutôt que d'afficher un cadre vide sur la page publique.")

# =========================================================================
#  PARTIE 3
# =========================================================================
partie('03', 'Les contenus')

rubrique('Actualités')
fiche([
    ('Alimente', "La page Actualités, la revue de presse, et l'aperçu sur l'accueil."),
    ('Obligatoire', "Titre, catégorie."),
    ('Pour publier', "Une date de publication. Et une source si la catégorie est « Presse »."),
    ('Images', "Une seule, choisie dans la médiathèque."),
])
p("Une actualité, c'est une parution, une dédicace, une conférence, un article de presse, "
  "un hommage. **Neuf catégories** sont proposées : Parution, Dédicace, Événement, "
  "Conférence, Presse, Reportage, Interview, Hommage, Archive retrouvée.")
p("**La catégorie « Presse » est particulière** : elle fait apparaître l'entrée dans la "
  "revue de presse, une page à part. C'est pourquoi le site exige alors le nom de l'organe "
  "dans le champ `Source` — un article sans son journal n'est pas vérifiable.")
p("**La date de publication classe la page.** Sans elle, une actualité se rangerait "
  "n'importe où dans la liste ; c'est pour cela qu'elle est exigée à la mise en ligne. "
  "Vous pouvez la postdater : l'entrée paraîtra avec cette date.")
p("Le `Chapô` est le paragraphe d'accroche affiché dans la liste, avant qu'on ouvre "
  "l'article. Quatre cents caractères au maximum, soit trois ou quatre phrases.")

rubrique('Événements')
fiche([
    ('Alimente', "L'agenda, et le bloc « à venir » de la page du livre."),
    ('Obligatoire', "Titre, date et heure de début."),
    ('Pour publier', "La ville."),
    ('Images', "Une seule."),
])
p("L'agenda sépare tout seul ce qui vient de ce qui a eu lieu : vous n'avez rien à "
  "archiver, la date s'en charge.")
p("La date de fin est facultative — un événement d'une soirée n'en a pas besoin. Si vous "
  "en mettez une, elle doit être postérieure au début : le site refuse l'inverse, qui "
  "donnerait une durée négative sur la page publique.")
p("**Une ville est exigée pour publier.** Un événement annoncé sans lieu n'aide personne à "
  "s'y rendre. Le champ `Lieu` porte le détail — la salle, l'institution ; le champ "
  "`Ville` porte la ville seule.")
p("Le `Lien d'inscription` est facultatif. S'il est rempli, un bouton d'inscription paraît "
  "sur la fiche. L'adresse doit commencer par `https://`.")

rubrique('Biographie')
fiche([
    ('Alimente', "La page Biographie, découpée en périodes ayant chacune sa propre adresse."),
    ('Obligatoire', "Titre."),
    ('Pour publier', "Les deux années de bornes, une source, et aucun chevauchement avec "
                     "une autre période publiée."),
    ('Images', "Plusieurs. La première cochée fait la vignette."),
])
p("La biographie n'est pas un long texte unique : elle est découpée en **périodes**. "
  "Chacune porte un titre, un récit, deux années de bornes, et vit à sa propre adresse — "
  "on peut donc envoyer le lien d'une seule période dans un dossier de presse.")
p("**Les bornes font tout le reste.** Une période qui connaît ses années sait "
  "automatiquement quels repères de la frise elle traverse et quelles pièces du fonds "
  "d'archives ont été produites pendant qu'elle durait. Vous ne rattachez rien à la main : "
  "les années suffisent.")
note("Trois conditions pour publier une période", [
    "**Les deux bornes** — sans elles, la période ne retrouve ni sa frise ni son fonds.",
    "**Une source** — c'est le seul endroit du site où l'on écrit la vie d'une personne "
    "réelle en continu, et une phrase non sourcée y passe pour un fait établi.",
    "**Aucun chevauchement** — deux périodes publiées ne peuvent pas partager une année. "
    "Sinon le même repère paraîtrait sous deux récits. Le site vous nomme la période qui gêne.",
])
p("Dans le récit, **une ligne vide sépare deux paragraphes**. C'est la seule mise en forme "
  "à connaître, et elle vaut pour tous les grands champs de texte du site.")
p("Cinq chapitres d'amorce vous attendent en brouillon. Ce sont d'anciens textes de "
  "calage, sans dates : ils sont là pour être remplacés, pas publiés tels quels.")

rubrique('Repères de la frise')
fiche([
    ('Alimente', "La frise chronologique de la biographie, et les quatre jalons de l'accueil."),
    ('Obligatoire', "Titre, année affichée, année de classement."),
    ('Pour publier', "Une source."),
    ('Images', "Aucune — une frise est du texte."),
])
p("Un repère est une date marquante : une naissance, une prise de fonction, une "
  "distinction. Il tient en une ligne et une courte notice.")
p("**Deux champs d'année, et ce n'est pas une erreur.**")
puces([
    "`Année affichée` — ce que le visiteur lit. Texte libre, parce qu'une date d'archive "
    "est souvent imprécise : « 1959 », « v. 1945 », « 1959-1960 », ou même « — ».",
    "`Année de classement` — un nombre, entre 1900 et 2100, qui range le repère au bon "
    "endroit sur la frise. Il n'apparaît nulle part. Pour une date incertaine, mettez votre "
    "meilleure estimation.",
])
p("**La case « Afficher sur la page d'accueil »** décide de ce qui remonte en page "
  "d'accueil, où la place est comptée : **les quatre premiers repères cochés** y "
  "paraissent, dans l'ordre chronologique. Cocher la case ne publie rien par elle-même — "
  "un repère en brouillon coché n'apparaît nulle part.")
p("La période à laquelle appartient un repère ne se saisit pas : elle se déduit de son "
  "année de classement.")

rubrique('Archives')
fiche([
    ('Alimente', "Le fonds d'archives : l'index, les catégories, et une page par pièce."),
    ('Obligatoire', "Titre, catégorie."),
    ('Pour publier', "Un crédit *ou* une source — au moins l'un des deux."),
    ('Fichiers', "Plusieurs, de toutes familles. Le premier coché fait la vignette."),
])
p("**Une archive est une notice, pas un fichier.** C'est la notion la plus importante de "
  "cet écran : vous ne cataloguez pas une photographie, vous décrivez une pièce — un "
  "événement, un document, un discours — à laquelle plusieurs fichiers peuvent se "
  "rattacher. Une même cérémonie peut porter quatre photographies, le programme en PDF et "
  "l'enregistrement du discours : c'est une seule archive.")
p("L'`Année` est facultative — une pièce est souvent mal datée — mais si vous la donnez, "
  "elle doit tomber entre 1900 et 2100 : c'est elle qui filtre et qui classe. Le champ "
  "`Date affichée`, lui, est du texte libre pour l'œil.")
p("Les champs `Personnes présentes` et `Mots-clés` nourrissent la recherche du site. "
  "Remplissez-les : c'est ce qui rendra le fonds consultable quand il comptera des "
  "centaines de pièces.")
note("Crédit ou source : l'un des deux, au minimum", [
    "Le `Crédit` nomme le photographe ou le détenteur des droits ; la `Source` nomme le "
    "fonds ou l'institution qui détient la pièce. Le site accepte de publier avec l'un ou "
    "l'autre — mais pas sans les deux.",
])

rubrique('Héritage')
fiche([
    ('Alimente', "La page Héritage, et une page par sujet."),
    ('Obligatoire', "Titre, rubrique."),
    ('Pour publier', "Une source."),
    ('Images', "Plusieurs. La première cochée fait la vignette."),
])
p("Héritage rassemble ce qui perpétue la mémoire : un pont, un boulevard, un buste, des "
  "décorations, des publications, une chanson. Cinq rubriques organisent l'index :")
puces([
    "**Lieux de mémoire** — un pont, un boulevard, un buste.",
    "**Hommages et commémorations**",
    "**Décorations et distinctions**",
    "**Livres et publications**",
    "**Musique et culture**",
])
p("**Seules les rubriques qui portent quelque chose apparaissent** sur le site. Vous "
  "pouvez donc ouvrir Héritage avec deux sujets : la page montre ce qui existe et se tait "
  "sur le reste.")
p("Le champ `Ordre` règle le rang à l'intérieur de la rubrique — le plus petit passe en "
  "premier.")
p("**La source est exigée pour publier**, et elle compte particulièrement ici : une liste "
  "de décorations se recopie de proche en proche avec ses erreurs. Publier la nôtre sans "
  "référence ajouterait une source de plus au malentendu.")

rubrique('Points de vente')
fiche([
    ('Alimente', "Le bloc « Où se procurer l'ouvrage », sur l'accueil et sur la page du livre."),
    ('Obligatoire', "La ville, et elle seule."),
    ('Pour publier', "Rien de plus."),
    ('Images', "Aucune."),
])
p("Une fiche par librairie, par enseigne ou par lieu de dédicace. La **ville** est le "
  "grand caractère affiché sur la carte ; l'enseigne et l'adresse viennent en dessous.")
p("Tout sauf la ville est facultatif, et **un champ vide ne s'affiche pas** : la carte se "
  "referme proprement sur ce qu'elle a. Vous pouvez donc créer « Bouaké » aujourd'hui et "
  "compléter l'enseigne la semaine prochaine.")
p("Si vous remplissez le champ `Site` — avec `https://` devant —, l'enseigne devient un "
  "lien vers celui-ci. Le téléphone, lui, devient un lien d'appel sur les mobiles.")
p("**Trois villes sont déjà en ligne** — Abidjan, Yamoussoukro, Paris — sans enseigne ni "
  "adresse. Elles affichent « Enseigne et adresse à renseigner », et l'écran vous le "
  "rappelle en haut de la liste tant que c'est le cas.")
p("Le nombre de points de vente est libre : la mise en page s'adapte toute seule, qu'il y "
  "en ait deux ou huit.")

# =========================================================================
#  PARTIE 4
# =========================================================================
partie('04', 'Ce qui vient des visiteurs')

p("Trois formulaires publics alimentent ces écrans. **Rien de ce qui arrive n'est jamais "
  "publié automatiquement** : tout attend votre décision. Ces écrans n'ont donc pas de "
  "bouton « créer » — on ne fabrique pas un témoignage, on décide de celui qu'on reçoit.",
  style='Lead')

rubrique('Témoignages')
p("Les visiteurs qui ont connu Philippe Grégoire Yacé peuvent déposer un souvenir. Chaque "
  "dépôt arrive **En attente**. Trois décisions vous sont offertes :")
puces([
    "**Publier** — le témoignage paraît sur la page Témoignages et peut remonter sur l'accueil.",
    "**Refuser** — il reste dans la liste, marqué comme refusé. Vous gardez la trace sans "
    "rien mettre en ligne.",
    "**Reprendre** — le remettre en attente, si vous changez d'avis.",
])
p("Vous pouvez **corriger un témoignage avant de le publier** — une faute de frappe, une "
  "ponctuation. Ouvrez sa fiche depuis la liste.")
p("Un visiteur ne peut pas envoyer plus de cinq témoignages par heure : c'est une "
  "protection contre les envois automatisés, invisible pour quelqu'un de bonne foi.")

rubrique('Contributions du public')
p("Le formulaire « Contribuez aux archives » permet à un visiteur d'envoyer des documents : "
  "une photographie de famille, une lettre, une coupure de presse. C'est le seul endroit du "
  "site qui reçoit des fichiers d'inconnus.")
note("Ces fichiers ne sont pas en ligne", [
    "Tant qu'une contribution n'est pas acceptée, ses fichiers sont retenus dans un espace "
    "fermé, inaccessible depuis l'extérieur. **La seule façon de les regarder est de passer "
    "par cet écran**, où vous êtes identifié. Aucun fichier envoyé par un inconnu ne devient "
    "visible sans votre accord.",
])
p("Ouvrez une contribution pour lire le message et examiner chaque fichier, puis :")
puces([
    "**Accepter** — les fichiers rejoignent la médiathèque, où vous les renseignez comme "
    "n'importe quel dépôt : titre, légende, crédit. Ils arrivent en brouillon ; c'est vous "
    "qui les publiez ensuite.",
    "**Refuser** — la contribution est marquée refusée.",
    "**Supprimer** — la contribution et ses fichiers disparaissent définitivement.",
])
p("Un visiteur peut joindre **cinq fichiers par envoi**, et trois envois par heure au "
  "maximum. Les formats et les poids acceptés sont exactement les mêmes que dans la "
  "médiathèque.")

rubrique('Messages')
p("Les messages du formulaire de contact arrivent ici plutôt que dans une boîte mail : "
  "rien ne se perd, et plusieurs personnes peuvent les traiter.")
p("Deux états seulement : **Nouveau** et **Traité**. Marquez un message comme traité une "
  "fois que vous y avez répondu — la réponse, elle, se fait depuis votre messagerie "
  "habituelle, en cliquant sur l'adresse du visiteur.")
p("Un message ne se modifie pas : il n'a pas vocation à être publié. Il se supprime, en "
  "revanche, quand il n'a plus d'utilité.")

# =========================================================================
#  PARTIE 5
# =========================================================================
partie('05', 'Le livre et la boutique')

rubrique('Paramètres')
p("Cet écran ne contient pas des contenus mais des **réglages** : des valeurs saisies une "
  "fois, qui alimentent plusieurs pages à la fois. Il est rangé en trois groupes.")

sous('La fiche technique de l’ouvrage')
p("Titre, auteur, éditeur, date de parution, nombre de pages, ISBN, prix, format. Ces "
  "valeurs paraissent sur la page du livre, et certaines sont reprises par les moteurs de "
  "recherche.")
puces([
    "**La date de parution est du texte libre** : « Mars 2026 » convient très bien si le "
    "jour n'est pas arrêté.",
    "**Le prix s'écrit en chiffres seuls**, sans espace ni devise : tapez `25000`, le site "
    "affiche « 25 000 F CFA ». C'est aussi ce prix qui sert à calculer les commandes.",
    "**L'ISBN** s'accepte avec ou sans tirets.",
])

sous('La préface')
p("Nom du préfacier, sa qualité, un extrait mis en exergue, et le texte complet. Une ligne "
  "vide sépare deux paragraphes.")
p("**La case « Mettre la préface en avant »** est un vrai levier éditorial : cochée, la "
  "préface remonte en tête de la page du livre et un bandeau paraît sur l'accueil. À cocher "
  "le jour où la préface est confirmée — un nom de préfacier suffit à l'annoncer, même "
  "avant que le texte n'arrive.")

sous('L’auteur du livre')
p("Nom, qualité, biographie. **La page de l'auteur se publie toute seule dès que le nom est "
  "renseigné**, et reste introuvable tant qu'il est vide — une fiche d'auteur sans auteur "
  "n'est pas une page.")
p("Attention à ne pas confondre l'auteur *du livre* avec son *sujet*.")

rubrique('Ouvrir les commandes')
p("La boutique est livrée **fermée**, et c'est délibéré : prendre une commande payable à "
  "la remise pour un ouvrage qui n'existe pas encore, c'est promettre une livraison qu'on "
  "ne peut pas tenir.")
note("Deux conditions, et il les faut toutes les deux", [
    "**La case « Ouvrir les commandes » cochée**, dans les paramètres.",
    "**Un prix saisi**, dans la fiche technique.",
    "L'une sans l'autre ne suffit pas : ouvrir sans prix afficherait « 0 F CFA » et "
    "enregistrerait des commandes gratuites. L'écran refuse la case tant que le prix est "
    "vide, et le dit.",
])
p("**Tant que la boutique est fermée**, la page Commander reste en ligne et annonce que les "
  "commandes ouvriront à la parution. C'est voulu : un bouton menant à une page introuvable "
  "fait croire à une panne.")
p("Deux réglages accompagnent l'ouverture :")
puces([
    "**Point de retrait** — l'adresse où l'on vient chercher un exemplaire, avec ses "
    "horaires. **Laissé vide, le retrait n'est pas proposé du tout** : offrir « retrait sur "
    "place » sans dire où est une promesse creuse.",
    "**Message affiché après commande** — ce que le client lit sous sa référence : le délai "
    "de rappel, le mode de paiement accepté à la remise, un numéro à joindre. Vide, le site "
    "s'en tient à sa formule standard.",
])
p("Le paiement se fait **à la livraison**. Aucun paiement en ligne n'est encaissé par le "
  "site aujourd'hui.")

rubrique('Zones de livraison')
p("*Réservé aux administrateurs.* Les frais de livraison se règlent zone par zone, dans une "
  "arborescence à trois niveaux : pays, ville, commune.")
p("**Les tarifs s'héritent**, et c'est ce qui rend le système tenable à la main. Laissez le "
  "champ des frais vide, et la zone prend le tarif de celle qui la contient. Vous posez donc "
  "« Côte d'Ivoire = 2 000 », puis seulement les exceptions : « Abidjan = 1 500 », "
  "« Cocody = 1 000 ».")
p("Sans cet héritage, il faudrait saisir un tarif pour chacune des treize communes du "
  "district, et pour chaque commune ajoutée ensuite.")
p("**Une zone dont ni elle ni aucune zone au-dessus ne porte de tarif n'est pas livrable** : "
  "le tunnel de commande ne la propose pas, plutôt que de laisser commander à zéro franc.")
p("Une zone qu'on ne dessert plus se **décoche** au lieu de se supprimer : les commandes "
  "anciennes gardent ainsi leur libellé.")

rubrique('Commandes')
p("*Réservé aux administrateurs.* Une commande naît du formulaire public. On ne la crée pas "
  "depuis l'administration, et **on ne la supprime jamais** : c'est une pièce comptable.")
p("Le parcours normal d'une commande en paiement à la livraison :")
etapes([
    "**Initiée** — le client vient de valider le formulaire.",
    "**Confirmée** — vous l'avez appelé, la commande est bonne.",
    "**Remise** — l'exemplaire est remis et encaissé. C'est le terme du parcours.",
])
p("À tout moment avant la remise, une commande peut être **annulée**. Les états sont "
  "verrouillés dans cet ordre : on ne peut pas sauter une étape ni revenir en arrière, ce "
  "qui évite qu'une commande remise repasse en attente.")
p("Chaque commande porte une **référence dictable au téléphone** — six signes du type "
  "`PGY-4F2K9A` — et un champ de **note** où consigner ce qui s'est dit : un rappel, un "
  "report, une adresse corrigée.")

# =========================================================================
#  PARTIE 6
# =========================================================================
partie('06', 'Réglages')

rubrique('Comptes')
p("*Réservé aux administrateurs.* Un compte par personne — jamais un compte partagé : les "
  "traces de modération et de remise perdraient le nom de qui a décidé.")
puces([
    "**Le mot de passe fait douze caractères au minimum.** Une phrase dont vous vous "
    "souvenez vaut mieux qu'un assemblage de symboles qu'il faudra noter quelque part.",
    "**Un compte se désactive, il ne se supprime pas.** C'est pourquoi il n'y a pas de "
    "bouton de suppression : effacer un compte effacerait la trace de ce qu'il a validé.",
    "Vous pouvez **réinitialiser le mot de passe** d'une personne qui a perdu le sien, "
    "depuis sa fiche.",
    "Le site refuse de désactiver ou de rétrograder **le dernier administrateur actif** : "
    "sans lui, plus personne ne pourrait rouvrir la porte.",
])

rubrique('Version anglaise')
p("Le site est bilingue. Les boutons, les menus et les libellés sont déjà traduits ; **ce "
  "que vous saisissez, non** — c'est l'objet de cet écran.")
p("**Un écran à part, et non un onglet sur chaque fiche.** Traduire n'est pas éditer : "
  "c'est souvent une autre personne, à un autre moment. Le traducteur n'a pas à traverser "
  "un formulaire de saisie, avec son sélecteur d'images et son bouton de publication, pour "
  "faire son travail.")
p("L'écran présente **le français à gauche et l'anglais à droite**, champ par champ. "
  "Choisissez une rubrique, puis une fiche.")
note("Un champ non traduit affiche le français", [
    "Il n'y a donc **aucun risque à traduire progressivement** : une page anglaise "
    "incomplète reste lisible, ce qui vaut mieux qu'une page à trous. Commencez par ce qui "
    "compte — les titres, les chapôs — et complétez au fil de l'eau.",
], sourd=True)
p("**Rien ne se publie depuis cet écran.** Poser une traduction ne change pas l'état d'une "
  "fiche et ne touche pas au français : c'est un écran qu'on peut confier sans risque.")
p("Tout ne se traduit pas, et c'est voulu. Une adresse web, une année, un crédit photo, un "
  "nom propre n'ont pas de version anglaise — seuls les champs qui le méritent vous sont "
  "proposés.")

# =========================================================================
#  PARTIE 7
# =========================================================================
partie('07', 'Aide-mémoire')

rubrique("Ce qui empêche de publier")
p("Le site refuse certaines mises en ligne. Ce n'est jamais arbitraire : chaque règle "
  "protège l'éditeur ou la lisibilité du site. Voici la liste complète.")
tableau(['Rubrique', 'Exigé pour publier', 'Pourquoi'], [
    ['Repères', 'Une source',
     "Rien n'est attribué à Yacé sans référence vérifiable."],
    ['Biographie', 'Les deux bornes, une source, aucun chevauchement',
     "Sans bornes, la période ne retrouve ni frise ni fonds. Deux périodes publiées ne "
     "peuvent pas partager une année."],
    ['Héritage', 'Une source',
     "Une liste de décorations se recopie avec ses erreurs."],
    ['Archives', 'Un crédit *ou* une source',
     "Une pièce publiée sans provenance expose l'éditeur."],
    ['Médiathèque', 'Un crédit',
     "Même raison : les fichiers viennent de familles et d'organes de presse."],
    ['Actualités', "Une date de publication, plus une source si catégorie « Presse »",
     "Sans date, l'entrée se range n'importe où. Un article sans son journal n'est pas "
     "vérifiable."],
    ['Événements', 'Une ville',
     "Un événement annoncé sans lieu n'aide personne."],
    ['Points de vente', 'Rien de plus que la ville',
     "Une ville seule est déjà une information utile."],
    ['Boutique', "La case cochée *et* un prix",
     "Ouvrir sans prix enregistrerait des commandes gratuites."],
], largeurs=[3.0, 5.2, 8.0])

rubrique("Fichiers en un coup d’œil")
tableau(['Question', 'Réponse'], [
    ["Formats d'image", 'JPEG, PNG, WebP — 8 Mo maximum'],
    ['Documents', 'PDF seulement — 30 Mo maximum'],
    ['Sons', 'MP3, M4A, OGG — 60 Mo maximum'],
    ['Vidéos', 'Non hébergées. Adresse YouTube collée dans la fiche'],
    ['Fichiers par dépôt', '20 au maximum'],
    ["Taille d'image maximale", '40 millions de pixels'],
    ['Taille conseillée', '1600 px sur le plus grand côté, JPEG qualité 80'],
    ['Poids visé par photo', 'entre 150 et 600 Ko'],
    ["Grande image d'accueil", '2000 × 2600 px, sujet à droite'],
    ['Couverture du livre', '1200 × 1550 px'],
    ['Portrait de biographie', '1400 × 1750 px'],
    ["Portrait de l'auteur", '1000 × 1250 px'],
    ['Extraits du livre', '1500 × 1000 px'],
    ['Archives et galerie', 'Aucune dimension imposée — sujet au centre'],
], largeurs=[5.4, 10.8])

rubrique("Si quelque chose ne marche pas")

sous("« J’ai enregistré, mais la page publique n’a pas changé »")
p("Neuf fois sur dix, le contenu est resté en **Brouillon**. Vérifiez la colonne Statut "
  "dans la liste. Sinon, rafraîchissez la page publique en forçant le rechargement.")

sous("« Mon image est refusée »")
p("Trois causes possibles, dans l'ordre de fréquence : elle dépasse 8 Mo ; ce n'est pas "
  "réellement un JPEG, un PNG ou un WebP malgré son nom ; elle dépasse 40 millions de "
  "pixels. Le message affiché vous dit laquelle.")

sous("« Le bouton de publication refuse et affiche un message »")
p("C'est une des règles du tableau ci-dessus. Le message nomme précisément ce qui manque. "
  "Ouvrez la fiche, complétez, réessayez.")

sous("« Je ne vois pas la rubrique décrite dans ce manuel »")
p("Votre compte est probablement un compte Éditeur : les commandes, les zones de livraison "
  "et les comptes sont réservés aux administrateurs. Elles sont retirées du menu, et non "
  "grisées.")

sous("« Je n’arrive plus à me connecter »")
p("Après cinq essais manqués, l'accès se bloque un quart d'heure. Attendez. Si vous avez "
  "perdu votre mot de passe, un administrateur peut le réinitialiser depuis l'écran Comptes.")

sous("« J’ai supprimé quelque chose par erreur »")
p("Il n'y a pas de corbeille : la suppression est définitive. Prenez l'habitude de "
  "**dépublier plutôt que de supprimer** — un contenu dépublié est invisible pour tout le "
  "monde et récupérable en un clic.")

# --- Pied de document ------------------------------------------------------
doc.add_paragraph()
fin = doc.add_paragraph()
bordure(fin._p, ['top'], FILET, taille=4)
fin.paragraph_format.space_before = Pt(10)
ecrire(fin, "Manuel d'administration du site **Philippe Grégoire Yacé — une destinée**. "
            "Rédigé le 8 septembre 2026, d'après le fonctionnement réel de l'administration. "
            "Les règles décrites ici sont celles que le site applique : quand il refuse une "
            "publication, c'est l'une d'elles qui parle.", taille=8, couleur=ENCRE3)

# --- Pied de page : numérotation ------------------------------------------
pied = doc.sections[0].footer.paragraphs[0]
pied.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = pied.add_run()
r.font.size = Pt(8)
r.font.name = TEXTE
r.font.color.rgb = ENCRE3
for instr in ('begin', 'PAGE', 'end'):
    e = OxmlElement('w:fldChar') if instr in ('begin', 'end') else OxmlElement('w:instrText')
    if instr in ('begin', 'end'):
        e.set(qn('w:fldCharType'), instr)
    else:
        e.set(qn('xml:space'), 'preserve')
        e.text = ' PAGE '
    r._element.append(e)

import os
sortie = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                      'Manuel-administration-site-Yace.docx')
doc.save(sortie)
print('Écrit :', sortie)
